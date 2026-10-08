#!/usr/bin/env python3
"""Heuristic static PPTX checker for astronomy talks.

Checks PowerPoint geometry/typography/alt text only. It cannot certify
scientific accuracy or visual legibility. See references/qa-rubric.md.

Usage:
    python check_pptx.py talk.pptx --output qa-static.md
    python check_pptx.py talk.pptx --json qa-static.json --strict
Dependencies: python-pptx (standard library otherwise)
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.enum.dml import MSO_COLOR_TYPE
    from pptx.util import Inches, Pt
except ImportError:
    print("Missing dependency: pip install python-pptx", file=sys.stderr)
    raise SystemExit(2)


def add(issues, severity, slide, code, detail):
    issues.append({"severity": severity, "slide": slide, "code": code, "detail": detail})


def norm_text(s):
    return " ".join((s or "").split())


def rgb_triplet(rgb):
    if rgb is None:
        return None
    try:
        return [int(rgb[0]), int(rgb[1]), int(rgb[2])]
    except Exception:
        return None


def luminance(rgb):
    parts = []
    for c in rgb:
        v = c / 255.0
        parts.append(v / 12.92 if v <= .04045 else ((v + .055) / 1.055) ** 2.4)
    return .2126 * parts[0] + .7152 * parts[1] + .0722 * parts[2]


def contrast(a, b):
    x, y = sorted([luminance(a), luminance(b)], reverse=True)
    return (x + .05) / (y + .05)


def slide_bg_rgb(slide):
    try:
        f = slide.background.fill
        if f.type is not None and f.fore_color.type == MSO_COLOR_TYPE.RGB:
            return rgb_triplet(f.fore_color.rgb)
    except Exception:
        pass
    return None  # Not safe to assume background when theme or shapes are used


def color_of_run(run):
    try:
        color = run.font.color
        if color and color.type == MSO_COLOR_TYPE.RGB:
            return rgb_triplet(color.rgb)
    except Exception:
        pass
    return None


def pic_alt_text(shape):
    try:
        nodes = shape._element.xpath(".//p:cNvPr")
        if not nodes:
            return ""
        return norm_text(nodes[0].get("descr", ""))
    except Exception:
        return ""


def run_font_sizes(shape):
    sizes = []
    unspecified = 0
    try:
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if not norm_text(run.text):
                    continue
                size = run.font.size or para.font.size
                if size is not None:
                    sizes.append(round(size.pt, 2))
                else:
                    unspecified += 1
    except Exception:
        pass
    return sizes, unspecified


def run_text_colors(shape):
    colors = []
    try:
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if norm_text(run.text):
                    c = color_of_run(run)
                    if c:
                        colors.append(c)
    except Exception:
        pass
    return colors


def check(pptx, *, min_font=18.0, max_chars=340):
    prs = Presentation(str(pptx))
    w, h = prs.slide_width, prs.slide_height
    aspect = w / h if h else 0
    issues = []
    stats = {"slides": len(prs.slides), "aspect_ratio": round(aspect, 4),
             "slide_width_in": round(w / Inches(1), 3),
             "slide_height_in": round(h / Inches(1), 3),
             "pictures": 0, "explicit_font_runs": 0, "text_chars": 0}
    if not (1.72 <= aspect <= 1.84):
        add(issues, "warning", "all", "aspect-ratio", f"Aspect is {aspect:.3f}; common astronomical meeting projector format is 16:9 (1.778). Check venue rules.")
    if not prs.slides:
        add(issues, "error", "all", "no-slides", "No slides found.")
    eps = Inches(.02)
    titles = Counter()
    for i, slide in enumerate(prs.slides, 1):
        texts, text_shape_count, pics = [], 0, 0
        bg = slide_bg_rgb(slide)
        for shape in slide.shapes:
            # Background overscan can be intentional, but tell user to inspect.
            if (shape.left < -eps or shape.top < -eps or
                    shape.left + shape.width > w + eps or
                    shape.top + shape.height > h + eps):
                add(issues, "warning", i, "out-of-bounds", f"'{shape.name}' lies partly outside slide bounds; verify intentional overscan.")
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pics += 1
                stats["pictures"] += 1
                if not pic_alt_text(shape):
                    add(issues, "warning", i, "missing-alt", f"Picture '{shape.name}' has no explicit description/alt text; mark decorative or add scientific meaning.")
            if not getattr(shape, "has_text_frame", False):
                continue
            raw = norm_text(shape.text)
            if not raw:
                continue
            text_shape_count += 1
            texts.append(raw)
            sizes, missing = run_font_sizes(shape)
            stats["explicit_font_runs"] += len(sizes)
            for s in sizes:
                # At the extreme bottom, smaller credits may be intentional.
                in_footer = (shape.top >= h - Inches(.64))
                if s < min_font and not in_footer:
                    add(issues, "warning", i, "small-font", f"'{shape.name}': explicit {s:g}pt < {min_font:g}pt (may be a figure inset, inspect at projector resolution).")
                    break
            if missing and not sizes:
                # Theme inheritance is common; not a defect by itself.
                pass
            # Check only when slide background is explicit, and text has no solid-filled container.
            if bg is not None:
                filled = False
                try:
                    filled = (shape.fill.type is not None)
                except Exception:
                    pass
                if not filled:
                    cols = run_text_colors(shape)
                    if cols:
                        worst = min(contrast(c, bg) for c in cols)
                        if worst < 4.5:
                            add(issues, "warning", i, "low-simple-contrast", f"'{shape.name}' ratio {worst:.2f}:1 against explicit slide background; check manually for overlays.")
        chars = sum(len(t) for t in texts)
        stats["text_chars"] += chars
        if chars > max_chars:
            add(issues, "warning", i, "text-heavy", f"{chars} text characters on slide; probably too dense for oral presentation.")
        if not texts and not pics:
            add(issues, "warning", i, "empty-slide", "No direct text or pictures found; could be intentional or use vectors/charts.")
        if pics > 5:
            add(issues, "info", i, "many-pictures", f"Contains {pics} pictures; verify panels and captions are readable.")
        if text_shape_count > 11:
            add(issues, "info", i, "many-text-boxes", f"{text_shape_count} nonempty text shapes; could be a crowded slide.")
        try:
            title = norm_text(slide.shapes.title.text) if slide.shapes.title is not None else ""
        except Exception:
            title = ""
        if not title and texts:
            # May be a deliberate custom layout, not always wrong.
            add(issues, "info", i, "no-title-placeholder", "No populated semantic title placeholder (custom headline can still be visually present).")
        if title:
            titles[title.lower()] += 1
    for k, n in titles.items():
        if n > 1:
            add(issues, "info", "all", "duplicate-semantic-title", f"Semantic title repeated {n} times: '{k[:75]}'.")
    counts = Counter(it["severity"] for it in issues)
    return {"file": str(pptx), "checked_at_utc": datetime.now(timezone.utc).isoformat(),
            "stats": stats, "counts": {x: counts[x] for x in ("error", "warning", "info")},
            "issues": issues,
            "disclaimer": "Static checks are heuristics. NOT a certification of scientific correctness, visual rendering, accessibility, or compatibility."}


def report_markdown(result):
    stats, counts = result["stats"], result["counts"]
    out = ["# Astronomy Talk PPTX — Static QA", "",
           f"- Source: `{result['file']}`", f"- Checked: `{result['checked_at_utc']}`",
           f"- Slides: **{stats['slides']}**; aspect: **{stats['aspect_ratio']}**",
           f"- Detected: **{counts['error']} errors**, **{counts['warning']} warnings**, **{counts['info']} info**", "",
           "## Issues", ""]
    if result["issues"]:
        out.append("| Severity | Slide | Code | Detail |")
        out.append("|---|---|---|---|")
        for it in result["issues"]:
            details = it["detail"].replace("|", "\\|").replace("\n", " ")
            out.append(f"| {it['severity']} | {it['slide']} | {it['code']} | {details} |")
    else:
        out.append("No static issues detected. This does **not** guarantee that the deck is ready to present.")
    out.extend(["", "## Mandatory manual gate", "",
                "- [ ] Every slide rendered and visually inspected at 16:9 full-screen.",
                "- [ ] Scientific claims/numbers/units/time systems/error bars traced to authentic sources.",
                "- [ ] Figure axis labels, legends, citations, credit, uncertainty and data filtering verified.",
                "- [ ] Slides timed by rehearsal; Q&A and transition excluded.",
                "- [ ] Tested playback in target conference software and projector.", "",
                f"> {result['disclaimer']}", ""])
    return "\n".join(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("pptx", type=Path, help="Path to .pptx deck")
    ap.add_argument("--output", type=Path, help="Save human-readable Markdown report")
    ap.add_argument("--json", type=Path, help="Save machine-readable JSON report")
    ap.add_argument("--strict", action="store_true", help="Exit 1 on warnings (errors always exit 1)")
    ap.add_argument("--min-body-font", type=float, default=18.0)
    ap.add_argument("--max-chars", type=int, default=340)
    a = ap.parse_args()
    if not a.pptx.exists():
        ap.error(f"File not found: {a.pptx}")
    try:
        result = check(a.pptx, min_font=a.min_body_font, max_chars=a.max_chars)
    except Exception as e:
        print(f"Cannot parse PPTX: {e}", file=sys.stderr)
        raise SystemExit(2)
    md = report_markdown(result)
    if a.output:
        a.output.parent.mkdir(parents=True, exist_ok=True)
        a.output.write_text(md, encoding="utf-8")
    if a.json:
        a.json.parent.mkdir(parents=True, exist_ok=True)
        a.json.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(md)
    return 1 if result["counts"]["error"] or (a.strict and result["counts"]["warning"]) else 0


if __name__ == "__main__":
    raise SystemExit(main())
